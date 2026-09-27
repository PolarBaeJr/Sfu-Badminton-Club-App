#!/usr/bin/env python3
"""Fold the CSVs auth-log-export.sh just wrote into one formatted workbook.

Called as: auth-log-xlsx.py <prefix>
where <prefix>-events.csv, -by-person.csv, -by-day.csv and -totals.csv exist.

Kept separate from the shell script so the CSV path has no Python dependency at
all. If openpyxl is missing the export still works, you just do not get the
formatted version.
"""
import csv
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F3864")
NOTE_FONT = Font(italic=True, size=9, color="595959")
BIG_FONT = Font(bold=True, size=20, color="1F3864")
STALE = PatternFill("solid", fgColor="FCE4E4")
THIN = Border(*[Side(style="thin", color="D9D9D9")] * 4)

# Sheet title, source suffix, explanatory note.
TABS = [
    ("By person", "by-person",
     "One row per account. Sorted by most recent activity. Rows shaded red have "
     "not been seen in over 30 days."),
    ("By day", "by-day",
     "Login counts per calendar day, Vancouver time."),
    ("Events", "events",
     "The raw log, newest first. Background token refreshes are excluded unless "
     "the export was run with --with-token-noise."),
]


def read_csv(path):
    if not os.path.exists(path):
        return [], []
    with open(path, newline="", encoding="utf-8") as fh:
        rows = list(csv.reader(fh))
    if not rows:
        return [], []
    return rows[0], rows[1:]


def looks_numeric(value):
    try:
        int(value)
        return True
    except (TypeError, ValueError):
        return False


def write_table(ws, headers, rows, start_row, stale_col=None):
    for i, name in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=i, value=name.replace("_", " "))
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)

    for r, row in enumerate(rows, start_row + 1):
        stale = False
        if stale_col is not None and stale_col < len(row) and looks_numeric(row[stale_col]):
            stale = int(row[stale_col]) > 30
        for c, value in enumerate(row, 1):
            cell = ws.cell(row=r, column=c,
                           value=int(value) if looks_numeric(value) else value)
            cell.border = THIN
            if stale:
                cell.fill = STALE

    # Size columns to the widest cell, capped so a long email does not push the
    # table off screen.
    for i, name in enumerate(headers, 1):
        widest = max([len(name)] + [len(r[i - 1]) for r in rows if i - 1 < len(r)] or [0])
        ws.column_dimensions[get_column_letter(i)].width = min(max(widest + 3, 11), 42)


def main():
    if len(sys.argv) != 2:
        print("usage: auth-log-xlsx.py <prefix>", file=sys.stderr)
        return 2
    prefix = sys.argv[1]

    _, total_rows = read_csv(f"{prefix}-totals.csv")
    totals = {k: v for k, v in total_rows}

    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Authentication log"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (f"SFU Badminton Club. Generated {totals.get('generated', '')} "
                f"Vancouver time, covering the last {totals.get('window_days', '?')} days.")
    ws["A2"].font = NOTE_FONT

    cards = [
        ("Accounts", totals.get("accounts_total", "?")),
        ("Ever signed in", totals.get("ever_signed_in", "?")),
        ("Active last 30 days", totals.get("signed_in_30d", "?")),
        ("Active last 7 days", totals.get("signed_in_7d", "?")),
    ]
    for i, (label, value) in enumerate(cards):
        col = 1 + i * 2
        ws.cell(row=4, column=col, value=label).font = Font(bold=True)
        cell = ws.cell(row=5, column=col,
                       value=int(value) if looks_numeric(value) else value)
        cell.font = BIG_FONT

    detail = [
        ("", ""),
        ("Login events in window", totals.get("login_events", "?")),
        ("Distinct people who logged in", totals.get("distinct_people_logged_in", "?")),
        ("Accounts that have never signed in", totals.get("never_signed_in", "?")),
        ("", ""),
        ("WHAT THIS DOES AND DOES NOT SHOW", ""),
        ("This is authentication only. It shows that somebody signed in, not which app they", ""),
        ("opened or whether they did anything once inside. auth.sessions does not separate", ""),
        ("the player app from the admin console.", ""),
        ("", ""),
        ("Background token refreshes are excluded. They are roughly 80% of the raw table and", ""),
        ("are a session staying alive, not a person logging in.", ""),
        ("", ""),
        ("'passkey or email code' in the method column is not missing data. Only OAuth logins", ""),
        ("record a provider, so anything else is correctly blank at the source.", ""),
    ]
    row = 7
    for label, value in detail:
        ws.cell(row=row, column=1, value=label)
        if label.isupper() and label:
            ws.cell(row=row, column=1).font = Font(bold=True, color="1F3864")
        if value:
            ws.cell(row=row, column=2,
                    value=int(value) if looks_numeric(value) else value).font = Font(bold=True)
        row += 1
    ws.column_dimensions["A"].width = 84
    ws.column_dimensions["B"].width = 14

    for title, suffix, note in TABS:
        headers, rows = read_csv(f"{prefix}-{suffix}.csv")
        if not headers:
            continue
        sheet = wb.create_sheet(title)
        sheet["A1"] = title
        sheet["A1"].font = TITLE_FONT
        sheet["A2"] = note
        sheet["A2"].font = NOTE_FONT
        stale_col = headers.index("days_since") if "days_since" in headers else None
        write_table(sheet, headers, rows, 4, stale_col=stale_col)

    out = f"{prefix}.xlsx"
    wb.save(out)
    print(f"  {'xlsx':<12} {len(wb.sheetnames):5d} tabs  {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
